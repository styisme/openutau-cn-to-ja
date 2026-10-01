# -*- coding: utf-8 -*-
"""
生成 ChineseToJapanesePhonemizer v2.1.6 安装与使用教程 (docx)
依赖：pip install python-docx
运行：python generate_manual.py
输出：ChineseToJapanesePhonemizer_v216_使用说明.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cjk_font(run, font_name="微软雅黑", size=None, bold=None,
                 italic=None, color=None):
    run.font.name = font_name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    sizes = {0: 22, 1: 16, 2: 13, 3: 12}
    colors = {0: (0x1F, 0x3A, 0x93), 1: (0x1F, 0x3A, 0x93),
              2: (0x30, 0x30, 0x30), 3: (0x30, 0x30, 0x30)}
    set_cjk_font(run, size=sizes.get(level, 12), bold=True,
                 color=colors.get(level, (0, 0, 0)))
    return h


def add_paragraph(doc, text, size=10.5, bold=False, italic=False,
                  color=None, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_cjk_font(run, size=size, bold=bold, italic=italic, color=color)
    if align is not None:
        p.alignment = align
    return p


def add_code_block(doc, code):
    p = doc.add_paragraph()
    run = p.add_run(code)
    set_cjk_font(run, font_name="Consolas", size=9.5)
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F2F2F2')
    pPr.append(shd)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    set_cjk_font(run, size=size)
    return p


def add_numbered(doc, text, size=10.5):
    p = doc.add_paragraph(style='List Number')
    run = p.add_run(text)
    set_cjk_font(run, size=size)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ''
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        set_cjk_font(run, size=10.5, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            run = p.add_run(str(val))
            set_cjk_font(run, size=10)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def main():
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # ---------- 封面 ----------
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ChineseToJapanesePhonemizer")
    set_cjk_font(run, size=24, bold=True, color=(0x1F, 0x3A, 0x93))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("v2.1.6 安装与使用教程")
    set_cjk_font(run, size=16, bold=True, color=(0x30, 0x30, 0x30))

    quote = doc.add_paragraph()
    quote.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = quote.add_run("让日语音源唱中文 —— 自动把中文歌词转换为日语发音的音素器插件")
    set_cjk_font(run, size=10.5, italic=True, color=(0x60, 0x60, 0x60))

    doc.add_paragraph()

    # ---------- 一、插件简介 ----------
    add_heading(doc, "一、插件简介", level=1)
    add_paragraph(doc, "ChineseToJapanesePhonemizer（简称 ZH to JA）是一个 OpenUtau 音素器插件，"
                       "用于让日语音源（CV / VCV / CVVC）直接演唱中文歌词，"
                       "无需手动把汉字翻译成日语假名。")
    add_paragraph(doc, "工作原理：")
    add_code_block(doc, "中文歌词 → 拼音 → 日语 mora 序列 → 匹配音源别名 → 渲染歌声")
    add_paragraph(doc, "你只需要在音符里输入中文汉字，音素器会自动完成后续所有转换。")

    # ---------- 二、系统要求 ----------
    add_heading(doc, "二、系统要求", level=1)
    add_table(
        doc,
        headers=["项目", "要求"],
        rows=[
            ["OpenUtau 版本", "v0.1.570 或更高（官方 win-x64 编译版）"],
            ["操作系统", "Windows 10 / 11（x64）"],
            ["音源", "任何日语音源（CV、VCV、CVVC 均支持）"],
            [".NET 运行时", ".NET 10 Desktop Runtime（最新版 OpenUtau 官方版已自带）"],
        ],
        col_widths=[4, 11],
    )
    note = doc.add_paragraph()
    run = note.add_run("提示：本插件针对 OpenUtau 官方的 .NET 10 版本编译。"
                       "如果你使用的是旧版 OpenUtau（.NET 8 时代），"
                       "插件将无法加载，请先升级 OpenUtau。")
    set_cjk_font(run, size=10, italic=True, color=(0xC0, 0x50, 0x00))

    # ---------- 三、安装步骤 ----------
    add_heading(doc, "三、安装步骤", level=1)

    add_heading(doc, "方式一：使用安装包（推荐）", level=2)
    add_paragraph(doc, "双击 MyZHtoJAPlugin-v2.1.6-Setup.exe，按提示下一步即可。")
    add_paragraph(doc, "安装程序会自动探测 OpenUtau 的 Plugins 目录（通常在 "
                       "Documents\\OpenUtau\\Plugins\\）。如果探测失败，"
                       "手动浏览到你自己的 Plugins 目录即可。")

    add_heading(doc, "方式二：手动安装", level=2)

    add_heading(doc, "第一步：找到 OpenUtau 数据目录", level=3)
    add_paragraph(doc, "新版 OpenUtau 的插件目录通常在用户文档下：")
    add_code_block(doc,
                   "C:\\Users\\<你的用户名>\\Documents\\OpenUtau\\\n"
                   "├── Plugins\\          ← 插件放这里\n"
                   "├── Singers\\\n"
                   "└── Cache\\")
    p = doc.add_paragraph()
    run = p.add_run("提示：")
    set_cjk_font(run, size=10.5, bold=True)
    run = p.add_run("旧版便携版可能位于安装目录下的 Plugins\\ 文件夹。"
                    "如果两个位置都有，都放一份最保险。")
    set_cjk_font(run, size=10.5)

    add_heading(doc, "第二步：关闭 OpenUtau", level=3)
    add_paragraph(doc, "务必完全退出 OpenUtau。检查方法：")
    add_numbered(doc, "关闭主窗口")
    add_numbered(doc, "打开任务管理器（Ctrl + Shift + Esc），确认没有 OpenUtau.exe 在运行")
    warn = doc.add_paragraph()
    run = warn.add_run("⚠ 如果 OpenUtau 正在运行，插件 DLL 会被锁定，复制会失败或 DLL 消失。")
    set_cjk_font(run, size=10, bold=True, color=(0xC0, 0x00, 0x00))

    add_heading(doc, "第三步：复制插件 DLL", level=3)
    add_paragraph(doc, "把 MyZHtoJAPlugin.dll 复制到 Plugins 文件夹：")
    add_code_block(doc, "C:\\Users\\<你的用户名>\\Documents\\OpenUtau\\Plugins\\MyZHtoJAPlugin.dll")
    add_paragraph(doc, "复制完成后，Plugins 目录里应该有：")
    add_code_block(doc,
                   "Plugins\\\n"
                   "├── MyZHtoJAPlugin.dll   ← 本插件\n"
                   "├── zh2ja.yaml           ← 配置文件（可选）\n"
                   "└── （其他原有插件）")

    add_heading(doc, "第四步：启动 OpenUtau 验证", level=3)
    add_numbered(doc, "双击 OpenUtau.exe 启动")
    add_numbered(doc, "新建一个音轨，或打开一个已有工程")
    add_numbered(doc, "在音轨设置面板里找到 音素器（Phonemizer）下拉菜单")
    add_numbered(doc, "列表中应该出现：\"ZH to JA\" 或 \"Chinese to Japanese Phonemizer\"")
    add_paragraph(doc, "看到它就说明安装成功。")

    # ---------- 四、使用方法 ----------
    add_heading(doc, "四、使用方法", level=1)

    add_heading(doc, "1. 选择音素器", level=2)
    add_paragraph(doc, "在音轨设置里，把音素器切换为 \"ZH to JA\"。")

    add_heading(doc, "2. 直接输入中文歌词", level=2)
    add_paragraph(doc, "在钢琴卷帘中，双击音符，直接输入中文汉字即可：")
    add_code_block(doc,
                   "音符1: 你\n"
                   "音符2: 好\n"
                   "音符3: 世\n"
                   "音符4: 界")
    p = doc.add_paragraph()
    run = p.add_run("规则：")
    set_cjk_font(run, size=10.5, bold=True)
    run = p.add_run("一个音符 = 一个汉字。不要把一句话塞进同一个音符。")
    set_cjk_font(run, size=10.5)

    add_heading(doc, "3. 特殊符号", level=2)
    add_table(
        doc,
        headers=["符号", "作用", "使用场景"],
        rows=[
            ["-", "起始标记", "放在乐句第一个音符，表示从静音开头"],
            ["R", "尾音收尾", "放在乐句最后一个音符后，让尾音自然衰减"],
            ["+", "延音", "让上一个字的发音延长到当前音符"],
            ["吸入", "换气声", "需要一个明显的呼吸点时使用"],
        ],
        col_widths=[2.5, 4, 8.5],
    )

    add_heading(doc, "4. 多音字处理", level=2)
    add_paragraph(doc, "音素器默认会按常见读音转换，遇到多音字（如\"长\"、\"行\"、\"重\"）读错时，"
                       "有两种解决方法：")
    p = doc.add_paragraph()
    run = p.add_run("方法一：直接写拼音")
    set_cjk_font(run, size=10.5, bold=True)
    add_paragraph(doc, "把歌词换成拼音，例如把\"长\"改成 chang 或 zhang。")
    p = doc.add_paragraph()
    run = p.add_run("方法二：使用音素提示")
    set_cjk_font(run, size=10.5, bold=True)
    add_paragraph(doc, "右键点击音符 → 选择\"音素提示\"（Phonetic Hint）→ 输入拼音，覆盖默认转换。")

    add_heading(doc, "5. 渲染与试听", level=2)
    add_paragraph(doc, "填完词后，点击 渲染（Render）按钮。如果一切正常，"
                       "就能听到日语音源唱出的中文。")

    add_heading(doc, "6. 更换 DLL / YAML 后必做", level=2)
    warn2 = doc.add_paragraph()
    run = warn2.add_run("⚠ 每次更换插件 DLL 或修改 zh2ja.yaml 后，请执行：")
    set_cjk_font(run, size=10.5, bold=True, color=(0xC0, 0x00, 0x00))
    add_numbered(doc, "完全退出 OpenUtau")
    add_numbered(doc, "复制新的 DLL 到 Plugins\\")
    add_numbered(doc, "清空 Documents\\OpenUtau\\Cache\\ 里的内容")
    add_numbered(doc, "启动 OpenUtau，重新渲染")
    add_paragraph(doc, "不清缓存会导致 OpenUtau 继续播放旧音素器生成的音频，"
                       "你改了参数也听不到效果。")

    # ---------- 五、v2.1.6 特性 ----------
    add_heading(doc, "五、v2.1.6 特性说明", level=1)
    add_paragraph(doc, "相比早期版本，主要有以下改进：")
    add_bullet(doc, "内置完整的 Romaji → 假名映射表，不依赖 WanaKanaNet，拗音（みゃ / きゅ / しょ 等）转换 100% 正确")
    add_bullet(doc, "完整拼音优先：400+ 常用拼音直接映射到日语谐音（如 hua → ふぁ、chi → ち）")
    add_bullet(doc, "多音节输入：一个音符里用空格分隔多个拼音")
    add_bullet(doc, "标点自动过滤、数字转拼音、儿化音合并")
    add_bullet(doc, "一/不变调（需要输入声调符号或数字后缀）")
    add_bullet(doc, "轻声压缩、叠字缩短、声调时长微调")
    add_bullet(doc, "支持 YAML 配置文件自定义映射")
    add_bullet(doc, "【v2.1.6 新增】鼻音韵尾三档处理：none / short / full")
    add_bullet(doc, "【v2.1.6 新增】short 模式采用 KOtoJA 风格「末尾固定毫秒」逻辑，把「ん」压缩到音符尾部")

    # ---------- 六、鼻音处理说明 ----------
    add_heading(doc, "六、鼻音韵尾（ん）处理说明", level=1)
    add_paragraph(doc, "中文的 an / en / ang / eng / in / ing / ong 等韵母，"
                       "转换到日语时会生成「ん」音素。不同音源的「ん」采样长度差异很大，"
                       "本插件提供三档处理模式：")

    add_table(
        doc,
        headers=["模式", "行为", "适用场景"],
        rows=[
            ["none", "完全不生成「ん」", "想要干净、不带鼻音，或音源的「ん」特别长"],
            ["short（默认）", "「ん」只占音符末尾固定毫秒", "想保留鼻音但不想拖长"],
            ["full", "保留完整「ん」采样", "想要明显的鼻音效果"],
        ],
        col_widths=[3, 6, 6],
    )

    add_paragraph(doc, "在 zh2ja.yaml 里配置：")
    add_code_block(doc,
                   "nasal_mode: short   # none / short / full\n"
                   "nasal_ms: 120       # short 模式下「ん」占音符末尾多少毫秒")

    p = doc.add_paragraph()
    run = p.add_run("⚠ 重要提醒：")
    set_cjk_font(run, size=10.5, bold=True, color=(0xC0, 0x00, 0x00))
    add_paragraph(doc, "「ん」的物理播放时长由音源的 oto.ini 决定，"
                       "音素器无法改变采样本身的长度。short 模式只是把「ん」的"
                       "起始位置推迟到音符末尾，因此：")
    add_bullet(doc, "普通音源（ん 采样 50ms 左右）→ short 模式效果明显")
    add_bullet(doc, "弱音源 / 气声音源（ん 采样 150ms+）→ short 模式效果有限，仍会偏长")

    add_paragraph(doc, "如果短也没用，请改用「分音符」方式，这是唯一 100% 精确的方案：")
    add_code_block(doc,
                   "音符1: a     ← 主元音\n"
                   "音符2: n     ← 单独的「ん」，音符长度决定实际时长")

    # ---------- 七、发音特点 ----------
    add_heading(doc, "七、发音特点说明", level=1)
    add_paragraph(doc, "由于中文和日语的音系差异，转换出来的发音会带有\"日式中文\"的特色，"
                       "属于正常现象：")
    add_table(
        doc,
        headers=["中文", "日语近似", "说明"],
        rows=[
            ["早 (zao)", "ざ お", "复合韵母拆成两个 mora"],
            ["喵 (miao)", "みゃ お", "拗音 + 单韵母"],
            ["穿 (chuan)", "ちゅ あ", "v2.1.6 中鼻音默认被压到末尾"],
            ["我 (wo)", "うぉ", "半元音近似"],
        ],
        col_widths=[3.5, 4, 7.5],
    )
    p = doc.add_paragraph()
    run = p.add_run("一句话总结：")
    set_cjk_font(run, size=10.5, bold=True)
    run = p.add_run("听起来像日本人唱中文，而不是标准普通话。"
                    "这是跨语言音源的固有限制，不是 Bug。")
    set_cjk_font(run, size=10.5)

    # ---------- 八、YAML 配置 ----------
    add_heading(doc, "八、YAML 配置说明", level=1)
    add_paragraph(doc, "插件启动时会自动读取以下位置的 zh2ja.yaml：")
    add_bullet(doc, "Documents\\OpenUtau\\Plugins\\zh2ja.yaml（全局配置）")
    add_bullet(doc, "音源目录\\zh2ja.yaml（该音源专属配置，优先级更高）")
    add_paragraph(doc, "没有这个文件时，插件使用内置默认值。示例配置：")
    add_code_block(doc,
                   "# zh2ja.yaml\n"
                   "nasal_mode: short       # none / short / full\n"
                   "nasal_ms: 120           # short 模式「ん」占末尾毫秒\n"
                   "disable_vcv: false      # true = 禁用 VCV 过渡\n"
                   "disable_light_tone: false\n"
                   "disable_tone_timing: false\n"
                   "disable_double_char: false\n"
                   "\n"
                   "# 自定义完整拼音映射（覆盖内置）\n"
                   "# full_pinyin:\n"
                   "#   hua: [fa]\n"
                   "#   chi: [chi]\n"
                   "\n"
                   "# 多音字修正\n"
                   "# polyphones:\n"
                   "#   \"chang da\": \"zhang da\"\n"
                   "\n"
                   "# 别名覆盖（弱音源想用短别名时）\n"
                   "# alias_overrides:\n"
                   "#   \"ん\": [\"ん -\", \"n -\", \"n\"]")

    # ---------- 九、FAQ ----------
    add_heading(doc, "九、常见问题（FAQ）", level=1)

    add_heading(doc, "Q1：安装后音素器列表里找不到 \"ZH to JA\"？", level=3)
    add_paragraph(doc, "排查顺序：")
    add_numbered(doc, "确认插件文件位置正确：应该在 Documents\\OpenUtau\\Plugins\\ 下。")
    add_numbered(doc, "确认文件名完整：不能是 MyZHtoJAPlugin.dll.txt 或带 (1) 后缀。")
    add_numbered(doc, "查看日志：菜单栏\"帮助\" → \"显示日志\"，搜索 ZH to JA。")
    add_numbered(doc, "重启 OpenUtau：有时候需要完全退出再启动。")

    add_heading(doc, "Q2：改了 YAML / 换了 DLL，效果没变化？", level=3)
    p = doc.add_paragraph()
    run = p.add_run("这是 OpenUtau 的缓存机制。")
    set_cjk_font(run, size=10.5, bold=True)
    add_paragraph(doc, "OpenUtau 会缓存已渲染的音频片段，音符内容不变就不会重新渲染。"
                       "换完插件后必须：")
    add_numbered(doc, "完全关闭 OpenUtau")
    add_numbered(doc, "清空 Documents\\OpenUtau\\Cache\\ 里的内容")
    add_numbered(doc, "重新启动并渲染")

    add_heading(doc, "Q3：渲染时报 Oto not found for \"ゃ\"？", level=3)
    p = doc.add_paragraph()
    run = p.add_run("这是正常的。")
    set_cjk_font(run, size=10.5, bold=True)
    add_paragraph(doc, "「ゃ」是日语拗音（みゃ / きゃ / しゃ 等）的小写后缀，"
                       "不是独立的 mora。OpenUtau 渲染时会做边界探测，"
                       "把「みゃ」拆成「み」和「ゃ」分别尝试查找 OTO。"
                       "该警告可以忽略，声音会正常播放。")

    add_heading(doc, "Q4：DLL 复制进去后消失了？", level=3)
    p = doc.add_paragraph()
    run = p.add_run("原因：")
    set_cjk_font(run, size=10.5, bold=True)
    add_bullet(doc, "OpenUtau 正在运行，锁定了文件。")
    add_bullet(doc, "你复制的是与官方内置插件同名的文件，被 OpenUtau 自动清理。")
    p = doc.add_paragraph()
    run = p.add_run("解决：")
    set_cjk_font(run, size=10.5, bold=True)
    add_numbered(doc, "完全关闭 OpenUtau（任务管理器确认）")
    add_numbered(doc, "使用独立文件名 MyZHtoJAPlugin.dll")
    add_numbered(doc, "重新复制，然后启动 OpenUtau")

    add_heading(doc, "Q5：某些字发音很奇怪？", level=3)
    add_paragraph(doc, "可以尝试：")
    add_numbered(doc, "换一种读法：用拼音代替汉字。")
    add_numbered(doc, "使用音素提示：右键音符 → 音素提示 → 手动指定发音。")
    add_numbered(doc, "换音源测试：不同音源的发音习惯差异很大。")
    add_numbered(doc, "反馈问题：把出问题的字和音源型号反馈给插件作者。")

    add_heading(doc, "Q6：ん 听起来太长怎么办？", level=3)
    add_paragraph(doc, "先确认 zh2ja.yaml 里的 nasal_mode 是 short 还是 none。"
                       "如果已经是 short 但还长，说明你的音源「ん」采样本身就很长"
                       "（常见于弱音源、气声音源）。此时：")
    add_bullet(doc, "方案一：改成 nasal_mode: none，直接省略鼻音")
    add_bullet(doc, "方案二：分两个音符写 a + n，用音符长度精确控制")
    add_bullet(doc, "方案三：调小 nasal_ms（如 60），能榨出一点效果")

    add_heading(doc, "Q7：能和官方内置的 \"JA VCV\" 或 \"JA CVVC\" 同时使用吗？", level=3)
    p = doc.add_paragraph()
    run = p.add_run("可以。")
    set_cjk_font(run, size=10.5, bold=True)
    run = p.add_run("本插件使用独立文件名 MyZHtoJAPlugin.dll，"
                    "不会与 OpenUtau 内置的音素器冲突。"
                    "你可以在不同音轨上选择不同音素器。")
    set_cjk_font(run, size=10.5)

    # ---------- 十、卸载 ----------
    add_heading(doc, "十、卸载方法", level=1)
    add_paragraph(doc, "如果使用安装包安装，直接到「控制面板 → 程序和功能」卸载即可。")
    add_paragraph(doc, "如果手动安装，只需要：")
    add_numbered(doc, "完全关闭 OpenUtau")
    add_numbered(doc, "删除 Documents\\OpenUtau\\Plugins\\MyZHtoJAPlugin.dll")
    add_numbered(doc, "重新启动 OpenUtau")
    p = doc.add_paragraph()
    run = p.add_run("不会影响 OpenUtau 的任何其他功能。")
    set_cjk_font(run, size=10.5, bold=True)

    # ---------- 十一、反馈与致谢 ----------
    add_heading(doc, "十一、反馈与致谢", level=1)
    add_paragraph(doc, "如果使用中遇到 Bug、有发音改进建议，"
                       "或者想反馈效果，欢迎联系作者。")
    add_table(
        doc,
        headers=["项目", "内容"],
        rows=[
            ["插件版本", "v2.1.6"],
            ["依赖", "OpenUtau v0.1.570+ / .NET 10 Desktop Runtime"],
            ["适用平台", "Windows x64"],
            ["安装方式", "双击安装包，或复制 DLL 到 Documents\\OpenUtau\\Plugins\\"],
        ],
        col_widths=[4, 11],
    )
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = end.add_run("感谢使用，祝调教愉快！ 🎤")
    set_cjk_font(run, size=12, bold=True, color=(0x1F, 0x3A, 0x93))

    output = "ChineseToJapanesePhonemizer_v216_使用说明.docx"
    doc.save(output)
    print(f"已生成：{output}")


if __name__ == "__main__":
    main()
